"""고객사 포털용 암호화 데이터 생성기 (목민 양식, 박스 단위 수불현황).
사용법: python build_portal_data.py <수불현황.xlsx> <고객사ID> <비밀번호> <표시이름> <출력폴더>
- 출력폴더/<ID 해시>.json 에 고객사 데이터를 비밀번호로 암호화해 저장한다.
- 비밀번호는 저장하지 않는다.
"""
import sys, os, json, base64, hashlib, secrets, datetime as dt
from zoneinfo import ZoneInfo
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

ITER = 310000
LOW = 5   # 품절 예정 기준(박스 미만)
TH = 2    # 곧 소진 기준(주)

def id_hash(cid):
    return hashlib.sha256(('freshline-portal:' + cid.strip()).encode('utf-8')).hexdigest()[:24]

def day(v):
    return (v.date() if isinstance(v, dt.datetime) else v).isoformat()

def num(v):
    try: return float(v or 0)
    except (TypeError, ValueError): return 0.0

def load_mokmin_box(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    st = wb['입출고현황']
    items, opening = [], {}
    for r in range(8, st.max_row + 1):
        n = st.cell(r, 2).value
        if not n or str(n).strip() == '합계': break
        items.append(str(n).strip()); opening[str(n).strip()] = num(st.cell(r, 5).value)
    rows = []
    for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True):
        if not r[0] or not r[1]: continue
        rows.append({'일자': day(r[0]), '구분': '입고', '점포': '', '품목': str(r[1]).strip(), '수량': num(r[3]), '비고': (str(r[4]).strip() if r[4] else '')})
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=7, values_only=True):
        if not r[0] or not r[3]: continue
        rows.append({'일자': day(r[0]), '구분': '출고', '점포': (str(r[2]).strip() if r[2] else ''), '품목': str(r[3]).strip(), '수량': num(r[5]), '비고': (str(r[6]).strip() if r[6] else '')})
    rows.sort(key=lambda x: (x['일자'], x['구분'] != '입고'))
    return items, opening, rows

def summarize(items, opening, rows):
    asof = max(r['일자'] for r in rows)
    end = dt.date.fromisoformat(asof); from4 = (end - dt.timedelta(days=27)).isoformat()
    stock = []
    for n in items:
        q = opening.get(n, 0) + sum(r['수량'] for r in rows if r['품목'] == n and r['구분'] == '입고') - sum(r['수량'] for r in rows if r['품목'] == n and r['구분'] == '출고')
        wk = sum(r['수량'] for r in rows if r['품목'] == n and r['구분'] == '출고' and from4 <= r['일자'] <= asof and '폐기' not in r['비고'] and '파손' not in r['비고']) / 4
        weeks = q / wk if wk > 0 else None
        status = '품절' if q <= 0 else ('곧 소진' if weeks is not None and weeks < TH else '여유')
        stock.append({'품목': n, '박스': q, '소진주': weeks, '상태': status})
    return asof, stock

def encrypt(payload, password):
    salt, iv = secrets.token_bytes(16), secrets.token_bytes(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(password.encode('utf-8'))
    ct = AESGCM(key).encrypt(iv, json.dumps(payload, ensure_ascii=False).encode('utf-8'), None)
    b = lambda x: base64.b64encode(x).decode()
    return {'v': 1, 'iter': ITER, 'salt': b(salt), 'iv': b(iv), 'ct': b(ct)}

if __name__ == '__main__':
    xlsx, cid, pw, name, outdir = sys.argv[1:6]
    items, opening, rows = load_mokmin_box(xlsx)
    asof, stock = summarize(items, opening, rows)
    checked = dt.datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M')
    payload = {'name': name, 'unit': '박스', 'asof': asof, 'checked': checked, 'low': LOW, 'stock': stock, 'rows': rows}
    os.makedirs(outdir, exist_ok=True)
    out = os.path.join(outdir, id_hash(cid) + '.json')
    json.dump(encrypt(payload, pw), open(out, 'w'))
    print(out, asof, [(s['품목'], s['박스'], s['상태']) for s in stock], len(rows))
