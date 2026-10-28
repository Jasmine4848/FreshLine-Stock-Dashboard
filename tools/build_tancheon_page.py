"""탄천갈매기 소스 재고 공개 페이지 생성기.
사용법: python build_tancheon_page.py <수불내역.xlsx> <출력 index.html>
"""
import sys, html, datetime as dt
import openpyxl
from zoneinfo import ZoneInfo

LOW = 5   # 품절 예정 기준(박스 미만)
TH = 2    # 상태 '곧 소진' 기준(주)

def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pi = wb['품목정보']
    items = []
    for r in range(5, pi.max_row + 1):
        n = pi.cell(r, 3).value
        if not n: continue
        items.append(dict(name=n, kg=pi.cell(r, 6).value, per=int(pi.cell(r, 7).value),
                          active=(pi.cell(r, 9).value or '') != '운영 중단'))
    def day(v): return v.date() if isinstance(v, dt.datetime) else v
    inb = [(day(r[0]), r[1], float(r[3] or 0)) for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True) if r[0]]
    out = []
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=6, values_only=True):
        if not r[0]: continue
        note = r[5] or ''
        kind = '출고' if '재배송' in note else ('폐기' if ('폐기' in note or '파손' in note) else '출고')
        out.append((day(r[0]), r[2], float(r[4] or 0), kind))
    return items, inb, out

def compute(items, inb, out):
    asof = max([d for d, *_ in inb] + [d for d, *_ in out])
    from4 = asof - dt.timedelta(days=27)
    rows = []
    for it in items:
        if not it['active']: continue
        n, per = it['name'], it['per']
        packs = sum(q for d, k, q in inb if k == n) - sum(q for d, k, q, _ in out if k == n)
        wk = sum(q for d, k, q, kind in out if k == n and kind == '출고' and from4 <= d <= asof) / 4
        weeks = packs / wk if wk > 0 else None
        status = '품절' if packs <= 0 else ('곧 소진' if weeks is not None and weeks < TH else ('여유' if wk > 0 else '최근 출고 없음'))
        rows.append(dict(name=n, spec=f"{it['kg']:g}kg × {per}팩", boxes=int(packs // per) if packs >= 0 else None,
                         weeks=weeks, status=status))
    return asof, rows

def render(asof, rows, checked):
    e = html.escape
    total = sum(r['boxes'] or 0 for r in rows)
    low = [r['name'] for r in rows if r['status'] != '품절' and r['boxes'] is not None and r['boxes'] < LOW]
    sold = [r['name'] for r in rows if r['status'] == '품절']
    def names(lst): return '<br>'.join(e(x) for x in lst) if lst else '없음'
    chip = {'품절': 'bad', '곧 소진': 'warn', '여유': 'ok', '최근 출고 없음': 'idle'}
    def weeks_txt(w): return '-' if w is None else ('1주 미만' if w < 1 else f'약 {int(w)}주')
    trs = '\n'.join(
        f'''        <tr><td><div class="name">{e(r['name'])}</div><div class="spec">{e(r['spec'])}</div></td>
          <td><span class="qty">{'-' if r['boxes'] is None else f"{r['boxes']:,}"}</span>{'' if r['boxes'] is None else '<small>박스</small>'}</td>
          <td class="days">{weeks_txt(r['weeks'])}</td>
          <td><span class="chip {chip[r['status']]}">{e(r['status'])}</span></td></tr>''' for r in rows)
    wd = '월화수목금토일'[asof.weekday()]
    return f'''<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<meta name="robots" content="noindex, nofollow">
<title>탄천갈매기 소스 재고</title>
<style>
  :root {{ --bg:#f5f6f4; --panel:#ffffff; --fg:#1d2321; --muted:#68716d; --line:#dfe3e0; --head:#eef1ef;
          --warn:#b45309; --bad:#c42b1c; --ok:#2f7d4f; --zebra:#f8f9f8; color-scheme: light;
          --font:"Malgun Gothic","맑은 고딕","Apple SD Gothic Neo","Noto Sans KR",sans-serif; }}
  @media (prefers-color-scheme: dark) {{ :root {{ --bg:#141816; --panel:#1c211f; --fg:#e8ece9; --muted:#9aa39e; --line:#2f3633; --head:#252b28;
          --warn:#f0a24a; --bad:#ff7a6b; --ok:#6cc58e; --zebra:#202623; color-scheme: dark; }} }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; background:var(--bg); color:var(--fg); font-family:var(--font); font-size:15px; line-height:1.5;
          padding: max(24px, env(safe-area-inset-top)) max(16px, env(safe-area-inset-right)) max(24px, env(safe-area-inset-bottom)) max(16px, env(safe-area-inset-left)); }}
  .wrap {{ max-width: 60rem; margin: 0 auto; display:grid; gap: 20px; }}
  h1 {{ margin:0; font-size: 1.75rem; font-weight:700; }}
  .lede {{ margin:4px 0 0; color:var(--muted); }}
  .lede b {{ color:var(--fg); }}
  .checked {{ margin:2px 0 0; font-size:12.5px; color:var(--muted); font-variant-numeric:tabular-nums; }}
  .tiles {{ display:grid; grid-template-columns: repeat(auto-fit, minmax(min(100%, 15rem), 1fr)); gap: 12px; }}
  .tile {{ --tone: var(--fg); background:var(--panel); border:1px solid var(--line); border-radius:10px; overflow:hidden; display:grid; grid-template-rows:auto 1fr; }}
  .tile-head {{ background:var(--head); border-bottom:1px solid var(--line); padding:8px 16px; display:flex; justify-content:space-between; align-items:baseline; gap:8px; flex-wrap:wrap; }}
  .tile-title {{ margin:0; font-weight:700; color:var(--tone); }}
  .tile-sub {{ margin:0; font-size:12.5px; color:var(--muted); }}
  .tile-body {{ padding:14px 16px 16px; display:flex; align-items:center; min-height:5.5rem; }}
  .value {{ font-size:1.9rem; font-weight:700; font-variant-numeric:tabular-nums; }}
  .value small {{ font-size:15px; font-weight:400; color:var(--muted); margin-left:4px; }}
  .names {{ font-weight:700; font-size:1.1rem; line-height:1.7; }}
  .alert:not(.on) .names {{ font-weight:400; color:var(--muted); font-size:15px; }}
  .alert.warn {{ --tone: var(--warn); }} .alert.bad {{ --tone: var(--bad); }}
  .alert.on {{ border-color: color-mix(in srgb, var(--tone) 45%, transparent); }}
  .alert.on .tile-head {{ background:var(--tone); border-bottom-color:var(--tone); }}
  .alert.on .tile-title {{ color:var(--panel); }}
  .alert.on .tile-sub {{ color: color-mix(in srgb, var(--panel) 85%, transparent); }}
  .alert.on .tile-body {{ background: color-mix(in srgb, var(--tone) 7%, var(--panel)); }}
  .alert.on .names {{ color:var(--tone); }}
  .card {{ background:var(--panel); border:1px solid var(--line); border-radius:10px; padding:20px; display:grid; gap:8px; min-width:0; }}
  .card h2 {{ margin:0; font-size:1.1rem; }}
  .note {{ margin:0; font-size:12.5px; color:var(--muted); }}
  .scroll {{ overflow-x:auto; }}
  table {{ width:100%; border-collapse:collapse; table-layout:fixed; font-variant-numeric:tabular-nums; }}
  col.c1 {{ width:34%; }} col.c2, col.c3, col.c4 {{ width:22%; }}
  th {{ font-size:12.5px; color:var(--muted); font-weight:700; padding:8px 12px; border-bottom:2px solid var(--line); text-align:center; white-space:nowrap; }}
  th:first-child {{ text-align:left; }}
  td {{ padding:10px 12px; border-bottom:1px solid var(--line); text-align:center; white-space:nowrap; }}
  td:first-child {{ text-align:left; }}
  tbody tr:nth-child(even) td {{ background:var(--zebra); }}
  tbody tr:last-child td {{ border-bottom:0; }}
  .name {{ font-weight:700; }} .spec {{ font-size:12.5px; color:var(--muted); }}
  .qty {{ font-weight:700; font-size:1.1rem; }} td small {{ font-size:12.5px; color:var(--muted); margin-left:3px; }}
  .chip {{ display:inline-flex; align-items:center; gap:6px; font-size:12.5px; font-weight:700; padding:2px 9px; border-radius:999px; background: color-mix(in srgb, currentColor 11%, transparent); }}
  .chip::before {{ content:""; width:7px; height:7px; border-radius:50%; background:currentColor; }}
  .chip.ok {{ color:var(--ok); }} .chip.warn {{ color:var(--warn); }} .chip.bad {{ color:var(--bad); }} .chip.idle {{ color:var(--muted); }}
  footer {{ display:flex; justify-content:center; padding-top:16px; margin-top:4px; border-top:1px solid var(--line); }}
  @media (max-width: 600px) {{
    table {{ table-layout:auto; }}
    col.c1, col.c2, col.c3, col.c4 {{ width:auto; }}
    .days {{ display:none; }}
    th, td {{ padding:10px 6px; }}
    .card {{ padding:16px 12px; }}
    .chip {{ white-space:normal; text-align:left; line-height:1.3; }}
  }}
  .brand {{ font-weight:700; font-size:1.5rem; letter-spacing:0.06em; }}
</style>
</head>
<body>
<div class="wrap">
  <header>
    <h1>탄천갈매기 소스 재고</h1>
    <p class="lede"><b>{asof:%Y-%m-%d} ({wd})</b> 기준 · 운영 중인 소스 <b>{len(rows)}</b>종</p>
    <p class="checked">마지막 확인: {checked:%Y-%m-%d %H:%M}</p>
  </header>
  <section class="tiles">
    <div class="tile"><div class="tile-head"><p class="tile-title">총 재고</p><p class="tile-sub">운영 품목 합계</p></div>
      <div class="tile-body"><span class="value">{total:,}<small>박스</small></span></div></div>
    <div class="tile alert warn{' on' if low else ''}"><div class="tile-head"><p class="tile-title">품절 예정</p><p class="tile-sub">재고 {LOW}박스 미만</p></div>
      <div class="tile-body"><div class="names">{names(low)}</div></div></div>
    <div class="tile alert bad{' on' if sold else ''}"><div class="tile-head"><p class="tile-title">품절</p><p class="tile-sub">재고 없음</p></div>
      <div class="tile-body"><div class="names">{names(sold)}</div></div></div>
  </section>
  <section class="card">
    <h2>품목별 현재고</h2>
    <p class="note days">예상 소진은 최근 4주 평균 출고량으로 계산한 대략적인 기간입니다.</p>
    <div class="scroll">
      <table>
        <colgroup><col class="c1"><col class="c2"><col class="c3"><col class="c4"></colgroup>
        <thead><tr><th>품목</th><th>총 보유 수량</th><th class="days">예상 소진</th><th>상태</th></tr></thead>
        <tbody>
{trs}
        </tbody>
      </table>
    </div>
  </section>
  <footer><span class="brand">(주)프레시라인</span></footer>
</div>
</body>
</html>
'''

if __name__ == '__main__':
    items, inb, out = load(sys.argv[1])
    asof, rows = compute(items, inb, out)
    checked = dt.datetime.now(ZoneInfo('Asia/Seoul'))
    open(sys.argv[2], 'w', encoding='utf-8').write(render(asof, rows, checked))
    print(asof, [(r['name'], r['boxes'], r['status']) for r in rows])
