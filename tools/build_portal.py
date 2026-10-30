"""(주)프레시라인 고객사 재고 조회 페이지 데이터 생성기.

사용법:
    python tools/build_portal.py <GSFM_정기업무 폴더> <설정 파일(json)> <출력 폴더: docs/portal/data>

- 설정 파일(고객사 목록·아이디·비밀번호·엑셀 위치)은 사용자 컴퓨터에만 두고 저장소에 넣지 않는다.
  형식: [{"id": 아이디, "pw": 비밀번호, "name": 화면 이름, "kind": 양식, "start": 시작일, "file": 엑셀 경로, ...}, ...]
- 고객사마다 데이터를 비밀번호로 암호화(AES-GCM, PBKDF2-SHA256)해 <아이디 해시>.json 으로 저장한다.
- 재고 계산은 각 엑셀의 입출고현황 수식과 같은 방식: 기초재고(시작일 기준) + 시작일 이후 입고 − 출고 (+ 조정).
"""
import sys, os, json, base64, hashlib, secrets, math, datetime as dt
from zoneinfo import ZoneInfo
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes

ITER = 310000
LOW_BOX = 5      # 품절 예정: 박스 환산 5박스 미만 (박스로 못 바꾸는 품목은 기본 단위 5 미만)
TH_WEEKS = 2     # (사용 안 함) 예전 출고 속도 기준
UNIT = {'BOX': '박스', 'PK': '팩', 'EA': '개'}


def id_hash(cid):
    return hashlib.sha256(('freshline-portal:' + cid.strip()).encode('utf-8')).hexdigest()[:24]

def d(v):
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    return None

def num(v):
    try: return float(v or 0)
    except (TypeError, ValueError): return 0.0

def s(v):
    return str(v).strip() if v is not None and str(v).strip() else ''

def unit_of(v, default):
    return UNIT.get(s(v).upper(), default)


# ---------- 양식별 읽기 ----------
def load_standard(wb, cfg):
    """표준 양식: 입출고현황(B 품목, [C 단위], 기초재고 열) / 입고내역 A~E / 출고내역 A 일자, B 브랜드, C 점포, D 품목, E 규격, F 수량, G 비고"""
    st = wb['입출고현황']
    name_col, unit_col, open_col = 2, cfg.get('unit_col'), cfg['open_col']
    items = []
    for r in range(8, st.max_row + 1):
        n = s(st.cell(r, name_col).value)
        if not n: break
        if n in ('합계', '소계'): continue
        u = unit_of(st.cell(r, unit_col).value, None) if unit_col else cfg.get('unit', '박스')
        if unit_col and u is None: continue          # 단위 없는 부자재 등은 제외
        items.append({'품목': n, '단위': u, '입수': None, '기초': num(st.cell(r, open_col).value)})
    rows = []
    for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True):
        if not d(r[0]) or not s(r[1]): continue
        rows.append({'일자': d(r[0]), '구분': '입고', '점포': '', '품목': s(r[1]), '수량': num(r[3]), '단위': unit_of(r[2], cfg.get('unit', '박스')), '비고': s(r[4]) if len(r) > 4 else ''})
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=7, values_only=True):
        if not d(r[0]) or not s(r[3]): continue
        rows.append({'일자': d(r[0]), '구분': '출고', '점포': s(r[2]), '품목': s(r[3]), '수량': num(r[5]), '단위': unit_of(r[4], cfg.get('unit', '박스')), '비고': s(r[6]) if len(r) > 6 else ''})
    return items, rows

def load_pack_info(wb, cfg):
    st, pi = wb['입출고현황'], wb['품목정보']
    per = {}
    for r in range(5, pi.max_row + 1):
        n = s(pi.cell(r, 3).value)
        if n: per[n] = (int(num(pi.cell(r, 7).value)) or None, s(pi.cell(r, 9).value) == '운영 중단')
    items = []
    for r in range(8, st.max_row + 1):
        n = s(st.cell(r, 2).value)
        if not n or n == '합계': break
        p, stopped = per.get(n, (None, False))
        if stopped: continue
        items.append({'품목': n, '단위': '팩', '입수': p, '기초': num(st.cell(r, 5).value)})
    rows = []
    for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True):
        if not d(r[0]) or not s(r[1]): continue
        rows.append({'일자': d(r[0]), '구분': '입고', '점포': '', '품목': s(r[1]), '수량': num(r[3]), '단위': '팩', '비고': s(r[4])})
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=6, values_only=True):
        if not d(r[0]) or not s(r[2]): continue
        rows.append({'일자': d(r[0]), '구분': '출고', '점포': s(r[1]), '품목': s(r[2]), '수량': num(r[4]), '단위': '팩', '비고': s(r[5])})
    return items, rows

def load_brand_pack(wb, cfg):
    st, pi = wb['입출고현황'], wb['품목정보']
    per = {}
    for r in range(6, pi.max_row + 1):
        n = s(pi.cell(r, 4).value)
        if not n or n == '매장명': break
        per[n] = int(num(pi.cell(r, 6).value)) or None
    items = []
    for r in range(8, st.max_row + 1):
        b, n = s(st.cell(r, 2).value), s(st.cell(r, 3).value)
        if b == '합계' or (not b and not n): break
        if not n or n == '소계': continue
        items.append({'품목': n, '브랜드': b, '단위': '팩', '입수': per.get(n), '기초': num(st.cell(r, 6).value)})
    rows = []
    for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True):
        if not d(r[0]) or not s(r[1]): continue
        rows.append({'일자': d(r[0]), '구분': '입고', '점포': '', '품목': s(r[1]), '수량': num(r[3]), '단위': '팩', '비고': s(r[4])})
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=6, values_only=True):
        if not d(r[0]) or not s(r[2]): continue
        rows.append({'일자': d(r[0]), '구분': '출고', '점포': s(r[1]), '품목': s(r[2]), '수량': num(r[4]), '단위': '팩', '비고': s(r[5])})
    return items, rows

def load_pack6(wb, cfg):
    st = wb['입출고현황']
    per = {}
    if '품목정보' in wb.sheetnames:                     # 박스입수: 품목정보 E열 ("6팩")
        pi = wb['품목정보']
        for r in range(5, pi.max_row + 1):
            n = s(pi.cell(r, 3).value)
            digits = ''.join(ch for ch in s(pi.cell(r, 5).value) if ch.isdigit())
            if n and digits: per[n] = int(digits)
    items = []
    for r in range(7, st.max_row + 1):
        n = s(st.cell(r, 2).value)
        if not n or n == '합계': break
        items.append({'품목': n, '단위': '팩', '입수': per.get(n, 6), '기초': num(st.cell(r, 8).value)})
    rows = []
    for r in wb['입고내역'].iter_rows(min_row=2, max_col=5, values_only=True):
        if not d(r[0]) or not s(r[1]): continue
        rows.append({'일자': d(r[0]), '구분': '입고', '점포': '', '품목': s(r[1]), '수량': num(r[3]), '단위': '팩', '비고': s(r[4])})
    for r in wb['출고내역'].iter_rows(min_row=2, max_col=7, values_only=True):
        if not d(r[0]) or not s(r[3]): continue
        rows.append({'일자': d(r[0]), '구분': '출고', '점포': s(r[2]), '품목': s(r[3]), '수량': num(r[5]), '단위': '팩', '비고': s(r[6])})
    return items, rows

def load_single_ledger(wb, cfg):
    st = wb['입출고현황']
    items = []
    for r in range(6, st.max_row + 1):
        n = s(st.cell(r, 3).value)
        if not n: continue
        if n == '합계': break
        p = st.cell(r, 4).value
        items.append({'품목': n, '단위': '개', '입수': int(p) if isinstance(p, (int, float)) and p else None, '기초': num(st.cell(r, 7).value)})
    rows = []
    for r in wb['수불내역'].iter_rows(min_row=2, max_col=6, values_only=True):
        if not d(r[0]) or not s(r[3]) or s(r[1]) not in ('입고', '출고', '조정'): continue
        rows.append({'일자': d(r[0]), '구분': s(r[1]), '점포': '', '품목': s(r[3]), '수량': num(r[4]), '단위': '개', '비고': s(r[5])})
    return items, rows

LOADERS = {'standard': load_standard, 'pack_info': load_pack_info, 'brand_pack': load_brand_pack, 'pack6': load_pack6, 'single_ledger': load_single_ledger}


# ---------- 계산 ----------
def build(cfg, base_dir, with_items=False):
    wb = openpyxl.load_workbook(os.path.join(base_dir, cfg['file']), data_only=True)
    items, rows = LOADERS[cfg['kind']](wb, cfg)
    nb = cfg.get('no_box')                     # 박스 환산 없이 낱개로만 표시: true(전체) 또는 품목 목록
    if nb:
        for it in items:
            if nb is True or it['품목'] in nb: it['입수'] = None
    start = cfg['start']
    rows.sort(key=lambda x: (x['일자'], {'입고': 0, '조정': 1, '출고': 2}[x['구분']]))
    asof = max(r['일자'] for r in rows) if rows else start
    end = dt.date.fromisoformat(asof)
    from4 = (end - dt.timedelta(days=13)).isoformat()
    waste = ('폐기', '파손')
    stock = []
    hide = set(cfg.get('hide_items') or [])   # 운영 중단 품목: 재고표에서만 빼고 입출고 내역은 그대로 둔다
    for it in items:
        if it['품목'] in hide: continue
        n = it['품목']
        mine = [r for r in rows if r['품목'] == n and r['일자'] >= start]
        q = it['기초'] + sum(r['수량'] for r in mine if r['구분'] == '입고') - sum(r['수량'] for r in mine if r['구분'] == '출고') + sum(r['수량'] for r in mine if r['구분'] == '조정')
        q = round(q, 2)
        wk = sum(r['수량'] for r in mine if r['구분'] == '출고' and from4 <= r['일자'] <= asof and not any(w in r['비고'] for w in waste)) / 2
        weeks = q / wk if wk > 0 else None
        boxes = q if it['단위'] == '박스' else (q // it['입수'] if it['입수'] else None)
        measure = boxes if boxes is not None else q      # 박스로 못 바꾸는 품목은 낱개 수량 그대로
        status = '품절' if q <= 0 else ('품절 예정' if measure < LOW_BOX else '여유')   # 5박스(낱개 품목은 5개) 미만
        stock.append({'품목': n, '브랜드': it.get('브랜드', ''), '수량': q, '단위': it['단위'], '입수': it['입수'], '소진주': weeks, '상태': status})
    if with_items: return asof, stock, rows, items
    return asof, stock, rows


# ---------- 배송료 ----------
FEE_SKIP = ('특판', '오배송', '오출고', '폐기')           # 비고에 있으면 배송료에서 제외
FEE_KEEP = ('재배송', '재출고', '교환')           # 단, 이 말이 함께 있으면 포함
def fee_period(cycle, asof):
    d = dt.date.fromisoformat(asof)
    if cycle == 'half':                         # 1~15일, 16일~말일
        last = (d.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
        return (d.replace(day=1), d.replace(day=15)) if d.day <= 15 else (d.replace(day=16), last)
    if cycle == 'week':                         # 월~토
        mon = d - dt.timedelta(days=d.weekday())
        return mon, mon + dt.timedelta(days=5)
    last = (d.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
    return d.replace(day=1), last

def fee_lines_sheet(f, base_dir):
    """배송현황 파일(합포장 등을 반영해 따로 관리하는 파일)에서 배송료 줄 만들기.
    box_list: '전체' 시트 B 일자, C 거래처, D 출고수량 → 수량 × 단가
    line_amount: '전체' 시트 B 출고유형, C 일자, D 거래처, E 품목, F 박스, J 총 합계금액"""
    src = f['source']
    wb = openpyxl.load_workbook(os.path.join(base_dir, src['file']), data_only=True, read_only=True)
    ws = wb['전체']
    out = []
    if src['kind'] == 'box_list':
        for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            if not isinstance(r[1], dt.datetime) or not isinstance(r[3], (int, float)): continue
            x = '모회사 출고' if '금산' in str(r[2] or '') else ''
            out.append({'d': r[1].date().isoformat(), 's': str(r[2] or ''), 'i': '', 'q': r[3], 'a': 0 if x else r[3] * f['price'], 'x': x})
    else:
        for r in ws.iter_rows(min_row=2, max_col=10, values_only=True):
            if not isinstance(r[2], dt.datetime) or not isinstance(r[9], (int, float)): continue
            x = '특판' if '특판' in str(r[1] or '') and not src.get('include_special') else ''
            out.append({'d': r[2].date().isoformat(), 's': str(r[3] or ''), 'i': str(r[4] or ''), 'q': r[5] or 0, 'a': 0 if x else r[9], 'x': x})
    wb.close()
    return out

def fee_lines_ledger(f, items, rows):
    per = {it['품목']: it['입수'] for it in items}
    out = []
    for r in rows:
        if r['구분'] != '출고': continue
        memo, store = r['비고'] or '', r['점포'] or ''
        fixed = '모회사 출고' if '금산' in store else ('재고 조정' if '조정' in store.replace(' ', '') else None)
        why = fixed or next((w for w in FEE_SKIP if w in memo), None)
        if why and not fixed and any(k in memo for k in FEE_KEEP): why = None
        if why and any(x.get('일자') == r['일자'] and x.get('품목') == r['품목'] for x in f.get('force_include', [])): why = None   # 사용자가 넣으라고 한 건
        if why:
            out.append({'d': r['일자'], 's': store, 'i': r['품목'], 'q': r['수량'], 'u': r['단위'], 'a': 0, 'x': why if fixed else memo}); continue
        if f.get('unit', '박스') == '박스':
            if r['단위'] == '박스': n = r['수량']
            elif per.get(r['품목']): n = r['수량'] / per[r['품목']]
            else: continue                          # 박스로 못 바꾸는 품목은 배송료 계산 안 함
            n = math.ceil(n - 1e-9)                # 0.6박스처럼 나누어떨어지지 않는 출고도 1박스로 계산
        else:
            n = r['수량']                       # 팩 단위 요금(열문집)
        out.append({'d': r['일자'], 's': store, 'i': r['품목'], 'q': n, 'a': n * f['price'], 'x': ''})
    return out

def fee_summary(cfg, items, rows, asof, base_dir=''):
    """배송료: 전체 기간의 줄을 담아 보내고(화면에서 기간별 조회), 현재 정산 기간 합계도 함께 계산."""
    f = cfg.get('fee')
    if not f: return None
    if f.get('source'):
        try: lines = fee_lines_sheet(f, base_dir)
        except Exception as e:
            print(cfg['id'], '배송현황 파일을 못 읽어 배송료 생략:', e); return None
    else:
        lines = fee_lines_ledger(f, items, rows)
    lines.sort(key=lambda x: x['d'])
    a, b = fee_period(f.get('cycle', 'month'), asof)
    cur = [x for x in lines if a.isoformat() <= x['d'] <= b.isoformat()]
    unit = f.get('unit', '박스') if not f.get('source') else '박스'
    return {'cycle': f.get('cycle', 'month'), 'price': f['price'], 'unit': unit, 'from': a.isoformat(), 'to': b.isoformat(),
            'qty': round(sum(x['q'] for x in cur if not x['x']), 2), 'total': round(sum(x['a'] for x in cur)),
            'unknown': [], 'lines': lines}

def encrypt(payload, password, cid=''):
    # salt 는 고객사마다 고정: 데이터가 매일 바뀌어도 같은 키가 나와서, 브라우저 탭에 로그인 상태를 유지할 수 있다
    salt = hashlib.sha256(('freshline-salt:' + cid.strip()).encode('utf-8')).digest()[:16] if cid else secrets.token_bytes(16)
    iv = secrets.token_bytes(12)
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=salt, iterations=ITER).derive(password.encode('utf-8'))
    ct = AESGCM(key).encrypt(iv, json.dumps(payload, ensure_ascii=False).encode('utf-8'), None)
    b = lambda x: base64.b64encode(x).decode()
    return {'v': 2, 'iter': ITER, 'salt': b(salt), 'iv': b(iv), 'ct': b(ct)}

def main(base_dir, cfg_file, out_dir, only=None):
    clients = json.load(open(cfg_file, encoding='utf-8'))
    pws = {c['id']: c.get('pw') for c in clients if c.get('pw')}
    checked = dt.datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y-%m-%d %H:%M')
    os.makedirs(out_dir, exist_ok=True)
    results = []
    built = {}
    for cfg in clients:
        if cfg.get('kind') == 'admin': continue
        if only and cfg['id'] not in only: continue
        if cfg['id'] not in pws:
            results.append((cfg['id'], '비밀번호 없음 - 건너뜀')); continue
        try:
            asof, stock, rows, items = build(cfg, base_dir, with_items=True)
        except Exception as e:  # 한 고객사 파일이 깨져도 나머지는 계속
            results.append((cfg['id'], f'실패: {e}')); continue
        payload = {'name': cfg['name'], 'asof': asof, 'checked': checked, 'low': LOW_BOX, 'stock': stock, 'rows': rows, 'fee': fee_summary(cfg, items, rows, asof, base_dir)}
        json.dump(encrypt(payload, pws[cfg['id']], cfg['id']), open(os.path.join(out_dir, id_hash(cfg['id']) + '.json'), 'w'))
        built[cfg['id']] = payload
        results.append((cfg['id'], asof, [(x['품목'], x['수량'], x['단위'], x['상태']) for x in stock], payload['fee'] and {k: payload['fee'][k] for k in ('from', 'to', 'qty', 'total', 'unknown')}))
    if built:                                   # 관리자 화면 데이터: 고객사 데이터가 하나라도 바뀌면 함께 다시 만든다
        for adm in [c for c in clients if c.get('kind') == 'admin' and c.get('pw')]:
            try:
                build_admin(adm, clients, pws, built, out_dir, checked)
                results.append((adm['id'], '관리자 화면 갱신'))
            except Exception as e:
                results.append((adm['id'], f'관리자 화면 실패: {e}'))
    for r in results: print(r)
    return results

def decrypt_file(path, password):
    blob = json.load(open(path))
    b = base64.b64decode
    key = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=b(blob['salt']), iterations=blob['iter']).derive(password.encode('utf-8'))
    return json.loads(AESGCM(key).decrypt(b(blob['iv']), b(blob['ct']), None))

def build_admin(adm, clients, pws, built, out_dir, checked):
    """관리자 계정: 모든 고객사 데이터를 한 파일에 담는다(이번에 안 만든 고객사는 기존 파일을 풀어서 쓴다). 시험용 계정(test)은 뺀다."""
    out = []
    for c in clients:
        if c.get('kind') == 'admin' or c.get('test') or c['id'] not in pws: continue
        pl = built.get(c['id'])
        path = os.path.join(out_dir, id_hash(c['id']) + '.json')
        if pl is None and os.path.exists(path):
            try: pl = decrypt_file(path, pws[c['id']])
            except Exception: pl = None
        if pl is None: continue
        pl = dict(pl, account=c['id'])
        out.append(pl)
    payload = {'admin': True, 'name': adm.get('name', '관리자'), 'checked': checked, 'low': LOW_BOX,
               'schedule': adm.get('schedule', []), 'clients': out}
    json.dump(encrypt(payload, adm['pw'], adm['id']), open(os.path.join(out_dir, id_hash(adm['id']) + '.json'), 'w'))

if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4:] or None)
