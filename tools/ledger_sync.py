"""통합 출고 시트 → 고객사 수불 엑셀 '출고내역' 자동 입력.

사용법:
    python tools/ledger_sync.py <GSFM_정기업무 폴더> <설정 json> <원본 xlsx> <상태 json> [--apply] [--today YYYY-MM-DD]

원칙 (보수적으로):
- 엑셀에 이미 있는 행은 절대 고치거나 지우지 않는다. 새 행만 맨 아래에 붙인다.
- 고객사별 '처리한 날짜'(상태 파일, 처음엔 엑셀 출고내역의 마지막 날짜) 다음 날부터만 넣는다.
- 품목·매장이 확실한 행(자동)만 넣는다. 추정·보류 행은 넣지 않고 보고한다.
- 넣으면 현재고가 마이너스가 되는 출고는 미출고일 수 있으므로 넣지 않고 보고한다.
- 이미 처리한 날짜에 원본 행이 뒤늦게 생기면 넣지 않고 보고한다.
- 엑셀 수정은 XML을 직접 고쳐 드롭다운(데이터 유효성)과 서식을 그대로 둔다.

설정 파일의 고객사 항목에 "sync" 가 있어야 대상이 된다:
    "sync": {"brands": [...], "tags": [...], "pk_as_box": false, "aliases": {"원본 품목": "엑셀 품목"}}
"""
import sys, os, re, json, zipfile, shutil, difflib, datetime as dt, collections, tempfile
from xml.sax.saxutils import escape
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_portal  # 재고 계산 재사용

# ---------------- 매칭 ----------------
def norm(s):
    s = str(s or '').lower()
    s = re.sub(r'\[[^\]]*\]', ' ', s)
    return re.sub(r'[\s()/·.,\-_~]', '', s)

def raw(s):
    return re.sub(r'[\s()\[\]/·.,\-_~]', '', str(s or '').lower())

MEAT = re.compile(r'(미국|호주|국내|캐나다|스페인|칠레|멕시코|뉴질랜드|브라질|아일랜드|덴마크|네덜란드)산|원물|슬라이스|포작업|분쇄육')

def sim(a, b):
    return difflib.SequenceMatcher(None, a, b).ratio()

class Client:
    def __init__(self, cid, brands, tags, items, stores, base_unit, per_box, aliases=None, pk_as_box=False):
        self.cid, self.items, self.stores = cid, items, stores
        self.brands = [raw(b) for b in brands]
        self.item_tags = [raw(t) for t in tags]
        self.base_unit, self.per_box, self.pk_as_box = base_unit, per_box, pk_as_box
        self.aliases = {norm(k): v for k, v in (aliases or {}).items()}
        self.strip = sorted({norm(x) for x in brands + tags if norm(x)}, key=len, reverse=True)
        self.keys = {it: self._strip(norm(it)) for it in items}

    def _strip(self, k):
        for b in self.strip:
            if k.startswith(b) and len(k) > len(b): return k[len(b):]
        return k

    def who(self, brand, store, item):
        rb, rs, ri = raw(brand), raw(store), raw(item)
        if any(t and t in ri for t in self.item_tags): return 1.0, '품목 브랜드 표시'
        if any(b and (rb == b or rb.startswith(b)) for b in self.brands): return 1.0, '브랜드'
        if any(b and rs.startswith(b) for b in self.brands): return 0.95, '점포명'
        best = max((sim(rs, raw(s)) for s in self.stores), default=0)
        if best >= 0.85: return 0.85, '점포명 유사'
        return 0, ''

    def canon_store(self, store):
        rs = raw(store)
        for s in self.stores:
            if raw(s) == rs: return s, 1.0
        best = max(self.stores, key=lambda s: sim(rs, raw(s)), default=None)
        if best and (sim(rs, raw(best)) >= 0.85 or (len(rs) >= 4 and raw(best).startswith(rs))):
            return best, sim(rs, raw(best))
        return str(store).strip(), 0.0

    def item(self, src):
        full = norm(src); n = self._strip(full)
        for key in (n, full):
            if key in self.aliases: return self.aliases[key], 1.0, '대응표'
        for it, k in self.keys.items():
            if n == k or full == norm(it): return it, 1.0, '일치'
        if MEAT.search(str(src)): return None, 0, '정육 상품'
        cands = [(it, k) for it, k in self.keys.items()
                 if k and any((k in x or x in k) and min(len(k), len(x)) >= 2 for x in (n, full))]
        if len(cands) == 1: return cands[0][0], 0.92, '포함'
        if cands: return max(cands, key=lambda c: len(c[1]))[0], 0.85, '포함(여러 후보)'
        if not self.keys: return None, 0, '없음'
        best = max(self.keys.items(), key=lambda kv: sim(n, kv[1]))
        sc = sim(n, best[1])
        if len(n) <= 3: sc = min(sc, 0.79)
        if sc >= 0.8: return best[0], sc, '비슷함'
        if sc >= 0.6: return best[0], sc, '약간 비슷'
        return None, sc, '없음'

    def qty(self, item, unit, q):
        u = str(unit or '').upper(); per = self.per_box.get(item)
        if 'BOX' in u:
            if self.base_unit == 'BOX': return q, None
            if per: return q * per, None
            return None, '박스입수 모름'
        if u in ('PK', 'EA'):
            if self.base_unit != 'BOX': return q, None
            if self.pk_as_box: return q, None
            return None, '원본은 팩, 엑셀은 박스 단위'
        return None, f'단위 {unit} 처리 불가'


def classify(rows, clients):
    out = []
    for r in rows:
        d, _, brand, store, item, unit, q = r
        cands = []
        for c in clients:
            w, why = c.who(brand, store, item)
            if w == 0: continue
            it, isc, iwhy = c.item(item)
            cands.append((c, w, why, it, isc, iwhy))
        if not cands: continue
        cands.sort(key=lambda x: (x[3] is not None and x[4] >= 0.8, x[1] * x[4]), reverse=True)
        c, w, why, it, isc, iwhy = cands[0]
        if it is None: continue                       # 위탁재고 품목이 아님(정육 등)
        qv, qerr = c.qty(it, unit, q)
        if isc >= 0.9 and w >= 0.95 and qerr is None: lvl = '자동'
        elif isc >= 0.8 and w >= 0.85 and qerr is None: lvl = '추정'
        elif isc < 0.6: continue
        else: lvl = '보류'
        cs, ssc = c.canon_store(store)
        if lvl == '자동' and ssc < 1.0 and w < 1.0: lvl = '추정'
        out.append({'고객사': c.cid, '일자': d, '점포': cs, '원본점포': store, '원본브랜드': brand, '원본품목': item,
                    '품목': it, '원본단위': unit, '원본수량': q, '수량': qv, '판정': lvl,
                    '근거': f'{why} / {iwhy}({isc:.2f})' + (f' / {qerr}' if qerr else '')})
    return out

# ---------------- 엑셀 읽기 ----------------
# 양식별 출고내역 열: 열 번호(1부터)
LAYOUT = {
    'standard':  {'date': 1, 'brand': 2, 'store': 3, 'item': 4, 'unit': 5, 'qty': 6, 'base': 'BOX'},
    'pack6':     {'date': 1, 'brand': 2, 'store': 3, 'item': 4, 'unit': 5, 'qty': 6, 'base': 'PK'},
    'pack_info': {'date': 1, 'store': 2, 'item': 3, 'unit': 4, 'qty': 5, 'base': 'PK'},
    'brand_pack': {'date': 1, 'store': 2, 'item': 3, 'box': 4, 'qty': 5, 'base': 'PK'},
}

def to_int(v):
    m = re.search(r'\d+', str(v or ''))
    return int(m.group()) if m else None

def read_client(cfg, path):
    wb = openpyxl.load_workbook(path, data_only=True)
    kind, L = cfg['kind'], LAYOUT[cfg['kind']]
    per, items = {}, []
    if kind == 'pack_info':
        pi = wb['품목정보']
        for r in range(5, pi.max_row + 1):
            n = pi.cell(r, 3).value
            if n: per[str(n).strip()] = to_int(pi.cell(r, 7).value)
        items = list(per)
    elif kind == 'brand_pack':
        pi = wb['품목정보']
        for r in range(6, pi.max_row + 1):
            n = pi.cell(r, 4).value
            if not n or n == '매장명': break
            per[str(n).strip()] = to_int(pi.cell(r, 6).value)
        items = list(per)
    elif kind == 'pack6':
        pi = wb['품목정보']
        for r in range(5, pi.max_row + 1):
            n = pi.cell(r, 3).value
            if n: per[str(n).strip()] = to_int(pi.cell(r, 5).value) or 6
        items = list(per)
    else:
        st = wb['입출고현황']
        for r in range(8, st.max_row + 1):
            n = st.cell(r, 2).value
            if not n or n == '합계': break
            if n != '소계': items.append(str(n).strip())
    ws = wb['출고내역']
    ledger, last_row = [], 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, L['date']).value
        if v is None: continue
        last_row = r
        if not isinstance(v, dt.datetime): continue
        g = lambda k: ws.cell(r, L[k]).value if k in L else None
        ledger.append({'일자': v.date(), '점포': str(g('store') or '').strip(), '품목': str(g('item') or '').strip(),
                       '수량': g('qty'), '브랜드': str(g('brand') or '').strip(), '단위': str(g('unit') or '').strip()})
    return items, per, ledger, last_row

# ---------------- 엑셀 쓰기 (XML) ----------------
def col_letter(c):
    s = ''
    while c: c, m = divmod(c - 1, 26); s = chr(65 + m) + s
    return s

def col_num(letters):
    n = 0
    for ch in letters: n = n * 26 + ord(ch) - 64
    return n

def sheet_path(z, name):
    wbx = z.read('xl/workbook.xml').decode('utf-8')
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(name), wbx).group(1)
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    tgt = re.search(r'<Relationship [^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels) or \
          re.search(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="%s"' % rid, rels)
    t = tgt.group(1)
    return t.lstrip('/') if t.startswith('/') else 'xl/' + t

ROW_RE = re.compile(r'<row [^>]*?r="(\d+)"[^>]*?(?:/>|>.*?</row>)', re.S)
CELL_RE = re.compile(r'<c [^>]*?r="([A-Z]+)(\d+)"[^>]*?(?:/>|>.*?</c>)', re.S)

def cell_xml(ref, style, value=None, formula=None):
    s = f' s="{style}"' if style is not None else ''
    if formula is not None:
        v = f'<v>{value}</v>' if value is not None else ''
        return f'<c r="{ref}"{s}><f>{escape(formula)}</f>{v}</c>'
    if value is None or value == '':
        return f'<c r="{ref}"{s}/>'
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"{s}><v>{value:g}</v></c>' if isinstance(value, float) else f'<c r="{ref}"{s}><v>{value}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'

def append_rows(path, out_path, sheet, start_row, new_rows):
    """new_rows: [{col번호: 값 또는 ('=f', 계산값)}] → start_row 부터 차례로 기록."""
    zin = zipfile.ZipFile(path)
    sp = sheet_path(zin, sheet)
    x = zin.read(sp).decode('utf-8')
    sd_start = x.index('<sheetData'); sd_end = x.index('</sheetData>')
    head, body, tail = x[:sd_start], x[sd_start:sd_end], x[sd_end:]
    rows = {int(m.group(1)): m.group(0) for m in ROW_RE.finditer(body)}
    # 서식 견본: 시작행 바로 윗줄(없으면 2행)
    tpl_row = rows.get(start_row - 1) or rows.get(2) or ''
    tpl = {col_num(m.group(1)): m.group(0) for m in CELL_RE.finditer(tpl_row)}
    def style_of(cxml):
        m = re.search(r'\ss="(\d+)"', cxml or '')
        return m.group(1) if m else None
    for i, vals in enumerate(new_rows):
        rn = start_row + i
        old = rows.get(rn)
        cells = {}
        if old:
            for m in CELL_RE.finditer(old): cells[col_num(m.group(1))] = m.group(0)
            for c, cx in cells.items():                       # 기존 값이 있으면 덮어쓰지 않도록 확인
                if c in vals and re.search(r'<v>|<is>', cx) and '<f>' not in cx:
                    raise RuntimeError(f'{sheet} {rn}행 {col_letter(c)}열에 이미 값이 있어 중단')
        for c in set(tpl) | set(vals):
            ref = f'{col_letter(c)}{rn}'
            base = cells.get(c) or tpl.get(c)
            st = style_of(base)
            f_old = re.search(r'<f>(.*?)</f>', cells.get(c) or '', re.S)
            if c in vals:
                v = vals[c]
                if isinstance(v, tuple):                       # (수식, 계산값)
                    cells[c] = cell_xml(ref, st, v[1], v[0])
                else:
                    cells[c] = cell_xml(ref, st, v)
            elif c not in cells:
                f_tpl = re.search(r'<f>(.*?)</f>', tpl.get(c) or '', re.S)
                if f_tpl:                                       # 윗줄 수식을 행 번호만 바꿔 복사
                    from xml.sax.saxutils import unescape
                    f = re.sub(r'(?<![A-Z$])([A-Z]{1,2})%d(?!\d)' % (rn - 1), lambda m: m.group(1) + str(rn), unescape(f_tpl.group(1)))
                    cells[c] = cell_xml(ref, st, None, f)
                else:
                    cells[c] = cell_xml(ref, st)
            elif f_old and c not in vals:
                pass
        attrs = ''
        if old:
            attrs = re.match(r'<row ([^>]*?)/?>', old).group(1)
            attrs = re.sub(r'\s*spans="[^"]*"', '', attrs)
        else:
            m = re.match(r'<row ([^>]*?)/?>', tpl_row)
            attrs = re.sub(r'\s*spans="[^"]*"', '', m.group(1)) if m else f'r="{rn}"'
            attrs = re.sub(r'r="\d+"', f'r="{rn}"', attrs)
        rows[rn] = f'<row {attrs}>' + ''.join(cells[c] for c in sorted(cells)) + '</row>'
    new_body = '<sheetData>' + ''.join(rows[k] for k in sorted(rows))
    x2 = head + new_body + tail
    last = max(rows)
    x2 = re.sub(r'<dimension ref="([A-Z]+\d+):([A-Z]+)(\d+)"/>',
                lambda m: f'<dimension ref="{m.group(1)}:{m.group(2)}{max(int(m.group(3)), last)}"/>', x2, count=1)
    tmp = out_path + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == sp: data = x2.encode('utf-8')
            elif info.filename == 'xl/workbook.xml':
                w = data.decode('utf-8')
                if '<calcPr' in w:
                    w = re.sub(r'<calcPr([^>]*?)\s*fullCalcOnLoad="[^"]*"', r'<calcPr\1', w)
                    w = w.replace('<calcPr', '<calcPr fullCalcOnLoad="1"', 1)
                else:
                    w = w.replace('</workbook>', '<calcPr fullCalcOnLoad="1"/></workbook>')
                data = w.encode('utf-8')
            zout.writestr(info, data)
    zin.close()
    os.replace(tmp, out_path)

# ---------------- 실행 ----------------
def excel_serial(d):
    return (d - dt.date(1899, 12, 30)).days

def key(d, store, item, q):
    return f'{d.isoformat()}|{raw(store)}|{item}|{float(q or 0):g}'

def load_source(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['통합출고내역']
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        if isinstance(r[0], dt.datetime) and r[4] and isinstance(r[6], (int, float)) and r[6] > 0:
            rows.append((r[0].date(),) + tuple(r[1:7]))
    return rows

def run(base_dir, cfg_file, src_file, state_file, apply=False, today=None, files=None):
    today = today or dt.date.today()
    clients_cfg = [c for c in json.load(open(cfg_file, encoding='utf-8')) if c.get('sync')]
    state = json.load(open(state_file, encoding='utf-8')) if os.path.exists(state_file) else {}
    src = [r for r in load_source(src_file) if r[0] <= today]
    src_max = max(r[0] for r in src)
    prepared = []
    for cfg in clients_cfg:
        path = (files or {}).get(cfg['id']) or os.path.join(base_dir, cfg['file'])
        items, per, ledger, last_row = read_client(cfg, path)
        sy = cfg['sync']; L = LAYOUT[cfg['kind']]
        stores = sorted({x['점포'] for x in ledger if x['점포']})
        cl = Client(cfg['id'], sy.get('brands', []), sy.get('tags', []), items, stores, L['base'], per,
                    sy.get('aliases'), sy.get('pk_as_box', False))
        prepared.append((cfg, path, cl, ledger, last_row, per))
    matched = classify(src, [p[2] for p in prepared])
    report, plans = {}, {}
    for cfg, path, cl, ledger, last_row, per in prepared:
        cid = cfg['id']; L = LAYOUT[cfg['kind']]
        mine = [m for m in matched if m['고객사'] == cid]
        st = state.get(cid)
        if not st:                                   # 처음: 엑셀 마지막 날짜까지 처리한 것으로 본다
            wm = max([x['일자'] for x in ledger], default=dt.date.fromisoformat(cfg['start']) - dt.timedelta(days=1))
            st = {'처리일': wm.isoformat(), '확인한원본': sorted({key(m['일자'], m['원본점포'], m['원본품목'], m['원본수량']) for m in mine if m['일자'] <= wm})}
        wm = dt.date.fromisoformat(st['처리일'])
        seen = set(st['확인한원본'])
        late = [m for m in mine if m['일자'] <= wm and key(m['일자'], m['원본점포'], m['원본품목'], m['원본수량']) not in seen]
        new = sorted([m for m in mine if m['일자'] > wm], key=lambda m: (m['일자'], m['점포']))
        have = collections.Counter(key(x['일자'], x['점포'], x['품목'], x['수량']) for x in ledger)
        # 현재고
        _, stock, _ = build_portal.build(dict(cfg, file=path), '')
        cur = {s['품목']: s['수량'] for s in stock}
        brand_of = collections.Counter((x['점포'], x['브랜드']) for x in ledger if x['브랜드'])
        unit_of = collections.Counter((x['품목'], x['단위']) for x in ledger if x['단위'])
        unit_any = collections.Counter(x['단위'] for x in ledger if x['단위'])
        add, hold = [], []
        for m in new:
            k = key(m['일자'], m['점포'], m['품목'], m['수량'] or 0)
            if have.get(k):
                have[k] -= 1; m['처리'] = '이미 있음'; continue
            if m['판정'] != '자동':
                m['처리'] = f"보류({m['판정']}): {m['근거']}"; hold.append(m); continue
            q = m['수량']
            if cfg['kind'] == 'brand_pack':
                p = per.get(m['품목'])
                if not p or q % p:
                    m['처리'] = '보류: 박스로 나누어떨어지지 않음'; hold.append(m); continue
            left = cur.get(m['품목'], 0) - q
            if left < 0:
                m['처리'] = f"보류: 재고 부족 (현재고 {cur.get(m['품목'], 0):g}, 출고 {q:g})"; hold.append(m); continue
            cur[m['품목']] = left
            m['처리'] = '입력'; add.append(m)
        rows = []
        for m in add:
            v = {L['date']: excel_serial(m['일자']), L['store']: m['점포'], L['item']: m['품목']}
            if 'brand' in L:
                b = [bb for (s, bb), _ in brand_of.most_common() if s == m['점포']]
                v[L['brand']] = b[0] if b else (m['원본브랜드'] or '')
            if 'unit' in L:
                u = [uu for (it, uu), _ in unit_of.most_common() if it == m['품목']]
                v[L['unit']] = u[0] if u else (unit_any.most_common(1)[0][0] if unit_any else m['원본단위'])
            if cfg['kind'] == 'brand_pack':
                v[L['box']] = int(m['수량'] // per[m['품목']])
                v[L['qty']] = ('IF(OR(C{r}="",D{r}=""),"",D{r}*VLOOKUP(C{r},품목정보!$D:$F,3,FALSE))', m['수량'])
            else:
                v[L['qty']] = m['수량']
            rows.append(v)
        plans[cid] = (path, last_row + 1, rows)
        new_wm = max(wm, src_max)
        seen |= {key(m['일자'], m['원본점포'], m['원본품목'], m['원본수량']) for m in mine if m['일자'] <= new_wm}
        report[cid] = {'처리일(전)': wm.isoformat(), '처리일(후)': new_wm.isoformat(), '마지막엑셀행': last_row,
                       '입력': add, '보류': hold, '이미있음': [m for m in new if m.get('처리') == '이미 있음'], '늦게들어온원본': late}
        state[cid] = {'처리일': new_wm.isoformat(), '확인한원본': sorted(seen)}
    if apply:
        for cid, (path, start, rows) in plans.items():
            if not rows: continue
            fixed = []
            for i, v in enumerate(rows):
                r = start + i
                fixed.append({c: ((val[0].replace('{r}', str(r)), val[1]) if isinstance(val, tuple) else val) for c, val in v.items()})
            append_rows(path, path, '출고내역', start, fixed)
        json.dump(state, open(state_file, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return report, src_max

def fmt(report, src_max, applied):
    out = [f'원본 마지막 날짜: {src_max}  ({"엑셀에 입력함" if applied else "미리보기 - 엑셀은 그대로"})']
    for cid, r in report.items():
        out.append(f"\n[{cid}] {r['처리일(전)']} 다음 날부터 → 입력 {len(r['입력'])}건 / 보류 {len(r['보류'])}건 / 이미 있음 {len(r['이미있음'])}건 / 늦게 들어온 원본 {len(r['늦게들어온원본'])}건")
        for m in r['입력']: out.append(f"  + {m['일자']} {m['점포']} | {m['품목']} {m['수량']:g}")
        for m in r['보류']: out.append(f"  ? {m['일자']} {m['원본점포']} | {m['원본품목']} {m['원본단위']} {m['원본수량']:g} → {m['품목']} | {m['처리']}")
        for m in r['늦게들어온원본']: out.append(f"  ! 늦게: {m['일자']} {m['원본점포']} | {m['원본품목']} {m['원본단위']} {m['원본수량']:g}")
    return '\n'.join(out)

if __name__ == '__main__':
    a = sys.argv[1:]
    apply = '--apply' in a
    today = None
    if '--today' in a: today = dt.date.fromisoformat(a[a.index('--today') + 1])
    pos = [x for i, x in enumerate(a) if not x.startswith('--') and (i == 0 or a[i - 1] != '--today')]
    rep, mx = run(pos[0], pos[1], pos[2], pos[3], apply, today)
    print(fmt(rep, mx, apply))
